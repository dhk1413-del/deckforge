"""「风格构建」离线回归：证据分层（主题 XML / 像素 / 结构）→ 令牌生成 → 预览 → 落盘 → 风格库可见。

零成本、不联网、不调模型、不依赖 LibreOffice：
  · pptx 走 python-pptx 现造 + 直接改写 theme1.xml（验证“色值只能来自证据”）
  · 图片走 Pillow 现画（验证图片输入只有像素证据）
  · 视觉模型/文本模型不在本测试覆盖范围（它们只写文字，且会花钱）
所有写盘动作都被重定向到临时目录，**绝不污染仓库 style-packs/**。
"""
import io
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)

from server.app import config, style_builder, stylepacks  # noqa: E402

PASS, FAIL = [], []


def check(name, ok, detail=""):
    (PASS if ok else FAIL).append(name)
    print("  %s %s%s" % ("✅" if ok else "❌", name, (" — " + detail) if detail else ""))


def _rem(node):
    """从 fontSize 令牌（如 '2.33rem'）里取出数值，用于比较层级大小。"""
    m = re.match(r"([\d.]+)rem", str((node or {}).get("fontSize") or ""))
    return float(m.group(1)) if m else 0.0


# ────────────────────────── 沙箱：把写盘全部重定向到临时目录 ──────────────────────────

TMP = tempfile.mkdtemp(prefix="sb-test-")
STYLE_ROOT = os.path.join(TMP, "style-packs")
os.makedirs(os.path.join(STYLE_ROOT, "styles"), exist_ok=True)
os.makedirs(os.path.join(STYLE_ROOT, "layouts"), exist_ok=True)
config.STYLE_DIR = STYLE_ROOT
stylepacks.STYLE_DIR = STYLE_ROOT
style_builder.BUILDER_ROOT = os.path.join(TMP, "builders")

ACCENT1 = "1E88E5"
ACCENT2 = "F4511E"


def make_pptx(path):
    """现造一份最小 pptx，并把 theme1.xml 的 accent1/accent2 改成已知值。"""
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(3):
        s = prs.slides.add_slide(blank)
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8), Inches(1.2))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "第 %d 页标题" % (i + 1)
        run.font.size = Pt(28 if i == 0 else 20)
        body = s.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(4), Inches(2))
        body.text_frame.paragraphs[0].add_run().text = "要点 A%d" % (i + 1)
        # 一个圆角矩形 + 一个椭圆：给几何统计喂数据
        from pptx.enum.shapes import MSO_SHAPE
        s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(6000000), Emu(1200000), Emu(1800000), Emu(900000))
        s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(8400000), Emu(1200000), Emu(500000), Emu(500000))
    prs.save(path)

    tmp_out = path + ".tmp.pptx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp_out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                txt = data.decode("utf-8")
                for slot, value in (("accent1", ACCENT1), ("accent2", ACCENT2)):
                    txt = re.sub(r"(<a:%s>\s*<a:srgbClr val=\")[0-9A-Fa-f]{6}(\")" % slot,
                                 r"\g<1>%s\g<2>" % value, txt)
                data = txt.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp_out, path)


def make_png(path, color=(30, 136, 229), size=(640, 360)):
    from PIL import Image
    im = Image.new("RGB", size, (255, 255, 255))
    for x in range(size[0]):
        for y in range(size[1]):
            if y > size[1] * 0.66:                     # 下方 1/3 一条色带
                im.putpixel((x, y), color)
    im.save(path)


def main():
    print("「风格构建」离线回归（%s）\n" % os.path.basename(TMP))
    pptx_path = os.path.join(TMP, "template.pptx")
    png_path = os.path.join(TMP, "reference.png")
    make_pptx(pptx_path)
    make_png(png_path)

    # PDF 全链路离线验证；禁止 OOXML 路径、模型与缓存重复渲染。
    import fitz
    from PIL import Image
    from unittest.mock import patch

    print("[PDF] 直接渲染、缓存、证据与格式优先级")
    with fitz.open() as doc:
        for n in range(2):
            page = doc.new_page(width=640, height=360)
            page.draw_rect(page.rect, fill=(0.9, 0.95, 1))
            page.insert_text((40, 60), "PDF reference %d" % (n + 1))
        pdf_raw = doc.tobytes()
    with patch.object(style_builder.source_pages, "build_pages", side_effect=AssertionError("LibreOffice forbidden")):
        uploaded = style_builder.save_uploads([("reference.pdf", pdf_raw), ("ignored.txt", b"x"), ("ignored.txt", b"x")])
        pdf_bid = uploaded["builder_id"]
        check("上传响应与状态警告均无重复且保留首次顺序",
              len(uploaded["warnings"]) == len(set(uploaded["warnings"])) and
              style_builder._read_state(pdf_bid)["warnings"] == uploaded["warnings"] and
              "ignored.txt" in uploaded["warnings"][0] and "PDF 输入" in uploaded["warnings"][1])
        check("PDF 上传 kind 与页数", uploaded["kind"] == "pdf" and uploaded["pages"] == 2)
        images = style_builder.page_images(pdf_bid)
        check("PDF 两页 PNG 宽度均为 900", len(images) == 2 and
              all(Image.open(path).width == 900 for _, path in images))
        check("PDF 落盘为 source.pdf", os.path.isfile(os.path.join(style_builder._builder_dir(pdf_bid), "source.pdf")))
        stamps = [os.stat(path).st_mtime_ns for _, path in images]
        with patch.object(fitz.Page, "get_pixmap", side_effect=AssertionError("cache must not render")):
            check("ensure_pages 幂等且页数不变", style_builder.ensure_pages(pdf_bid) == images and
                  style_builder.page_count(pdf_bid) == 2 and
                  stamps == [os.stat(path).st_mtime_ns for _, path in images])
            check("start_pages PDF 缓存复用", style_builder.start_pages(pdf_bid) is False)
        with patch.object(style_builder, "_compose_prose", return_value=({}, {"ok": False, "reason": "offline_test"})), \
             patch.object(style_builder, "_theme_evidence", side_effect=AssertionError("no OOXML")), \
             patch.object(style_builder, "_structure_evidence", side_effect=AssertionError("no OOXML")), \
             patch.object(style_builder, "_vision_describe", side_effect=AssertionError("no model")):
            result = style_builder.analyze(pdf_bid, use_vision=False)
            first_warnings = list(result["warnings"])
            check("analyze 警告无重复且上传警告最先出现",
                  len(first_warnings) == len(set(first_warnings)) and
                  first_warnings[:len(uploaded["warnings"])] == uploaded["warnings"])
            check("响应与持久化警告一致且无重复",
                  style_builder.response(pdf_bid)["result"]["warnings"] == first_warnings and
                  style_builder._read_state(pdf_bid)["warnings"] == first_warnings and
                  style_builder._read_state(pdf_bid)["result"]["warnings"] == first_warnings)
            rerun = style_builder.analyze(pdf_bid, use_vision=False)
            check("重跑 analyze 警告不叠加且顺序不变",
                  rerun["warnings"] == first_warnings and
                  len(rerun["warnings"]) == len(set(rerun["warnings"])) and
                  style_builder.response(pdf_bid)["result"]["warnings"] == first_warnings)

        warning = "；".join(result["warnings"])
        check("PDF 如实说明主题色板与嵌入字体边界", "PDF" in warning and "OOXML 主题色板" in warning and "嵌入字体名" in warning)
        check("PDF 不编造 OOXML 色槽或 fontScheme", "theme" not in result["evidence"])
        check("无嵌入字体不填占位证据", "pdf_fonts" not in result["evidence"])
        check("PDF evidence 无空 theme/structure", all(k not in result["evidence"] for k in ("theme", "structure")))
        # 模拟修复前磁盘状态；GET 必须迁移并写回，保留跨列表首次出现的顺序。
        legacy = style_builder._read_state(pdf_bid)
        legacy["warnings"] = ["先出现", "后出现", "先出现"]
        legacy["result"]["warnings"] = ["后出现", "新增", "先出现"]
        legacy["result"]["evidence"].update(theme=dict(file="", name="", slots={}, fonts={}), structure={})
        with open(style_builder._state_path(pdf_bid), "w", encoding="utf-8") as f:
            json.dump(legacy, f, ensure_ascii=False)
        migrated = style_builder.response(pdf_bid)["result"]
        expected = ["先出现", "后出现", "新增"]
        check("历史警告按首次出现顺序去重并写回",
              migrated["warnings"] == expected and
              style_builder._read_state(pdf_bid)["warnings"] == expected and
              style_builder._read_state(pdf_bid)["result"]["warnings"] == expected)
        check("历史 PDF 空证据清理并写回", all(k not in style_builder._read_state(pdf_bid)["result"]["evidence"]
                                                for k in ("theme", "structure")))
        for kind in ("pdf", "image"):
            state = dict(kind=kind, warnings=["首次", "第二", "首次"], result=dict(
                warnings=["第二", "第三"], evidence=dict(theme={}, structure={})))
            style_builder._write_state("normalize_" + kind, state)
            check(kind + " 写状态警告有序去重且无空证据",
                  state["warnings"] == ["首次", "第二", "第三"] and state["result"]["evidence"] == {})
            state["result"]["evidence"] = dict(theme=dict(slots={"accent1": "#1E88E5"}, source="reference_xml"),
                                                structure=dict(shapes={"rect": 1}, source="reference_shapes"))
            style_builder._write_state("normalize_" + kind, state)
            check(kind + " 真实证据保留并说明来源",
                  len(state["result"]["evidence"]) == 2 and
                  any("reference_xml" in w for w in state["warnings"]) and
                  any("reference_shapes" in w for w in state["warnings"]))
        # 使用本地字体构造真实嵌入字体，验证证据来源。
        font_path = fitz.Font("helv").buffer
        with fitz.open() as embedded:
            page = embedded.new_page()
            page.insert_font(fontname="ReferenceFont", fontbuffer=font_path)
            page.insert_text((40, 60), "Embedded reference", fontname="ReferenceFont")
            font_pdf = style_builder.save_uploads([("embedded.pdf", embedded.tobytes())])
        font_evidence = style_builder._pdf_fonts(style_builder._pdf_path(style_builder._read_state(font_pdf["builder_id"])))
        check("提取真实嵌入字体且标注 pdf_embedded", font_evidence.get("source") == "pdf_embedded" and bool(font_evidence.get("names")))
    with open(pptx_path, "rb") as f:
        mixed = style_builder.save_uploads([("template.pptx", f.read()), ("reference.pdf", pdf_raw)])
    check("PPTX 优先于 PDF，丢弃并警告", mixed["kind"] == "pptx" and len(mixed["files"]) == 1 and "reference.pdf" in "；".join(mixed["warnings"]))
    with open(png_path, "rb") as f:
        mixed = style_builder.save_uploads([("reference.pdf", pdf_raw), ("reference.png", f.read())])
    check("PDF 优先于图片，丢弃并警告", mixed["kind"] == "pdf" and mixed["pages"] == 2 and len(mixed["files"]) == 1 and "reference.png" in "；".join(mixed["warnings"]))

    # ── ① 主题证据 ──
    print("[1] 主题 XML 证据（色值的权威来源）")
    theme = style_builder._theme_evidence(pptx_path)
    slots = theme.get("slots") or {}
    check("读到 theme1.xml 的色槽", bool(slots), "槽位 %d 个" % len(slots))
    check("accent1 与注入值一致", slots.get("accent1", "").upper() == "#" + ACCENT1,
          "得到 %s" % slots.get("accent1"))
    check("accent2 与注入值一致", slots.get("accent2", "").upper() == "#" + ACCENT2,
          "得到 %s" % slots.get("accent2"))
    fonts = theme.get("fonts") or {}
    check("读到字体方案", bool(fonts.get("majorFont") or fonts.get("minorFont")), str(fonts)[:60])

    # ── ② 结构证据 ──
    print("\n[2] 结构证据（字号/几何/圆角）")
    struct = style_builder._structure_evidence(pptx_path)
    sizes = struct.get("sizes_pt") or []
    check("统计到字号", bool(sizes), str(sizes[:4])[:80])
    shapes = struct.get("shapes") or {}
    check("统计到几何类型", bool(shapes), str(dict(list(shapes.items())[:4]))[:70])
    check("几何里含圆角矩形", "roundRect" in shapes or "roundedRectangle" in str(shapes), str(list(shapes)[:6]))

    # ── ③ 像素证据（图片输入的唯一证据源） ──
    print("\n[3] 像素证据")
    pages = [(1, png_path)]
    pixels = style_builder._pixel_evidence(pages)
    check("像素统计产出色表", bool(pixels.get("colors")), "%d 个饱和色" % len(pixels.get("colors") or {}))
    check("采样页记录正确", pixels.get("sampled_pages") == [1], str(pixels.get("sampled_pages")))
    rec = style_builder._find_pixel_color(pixels, "#" + ACCENT1, tol=40)
    check("能命中注入的主色（容差内）", rec is not None,
          "命中 %s" % (rec or {}).get("hex"))
    check("找不到的颜色返回 None（不编造）",
          style_builder._find_pixel_color(pixels, "#00FF00", tol=6) is None)

    # ── ④ 合成色板：每个值都要能说出出处 ──
    print("\n[4] 色板合成（证据 → 令牌）")
    warnings = []
    palette, evidence = style_builder.build_palette(theme, pixels, warnings)
    check("accent 取主题 accent1", str(palette.get("accent", "")).upper() == "#" + ACCENT1,
          "得到 %s" % palette.get("accent"))
    check("accent-2 有值", bool(palette.get("accent-2")), str(palette.get("accent-2")))
    check("每条证据都带 source", all(e.get("source") for e in evidence), "%d 条" % len(evidence))
    check("证据来源可追溯到 theme / pixels / derived",
          all(str(e["source"]).split(":")[0] in ("theme", "pixels", "derived") for e in evidence),
          "、".join(sorted({str(e["source"]).split(":")[0] for e in evidence})))
    check("色值格式合法", all(re.match(r"^#[0-9A-Fa-f]{6}$", v) for v in palette.values() if isinstance(v, str)))

    # ── ⑤ 令牌与 style.md ──
    print("\n[5] 令牌与 style.md")
    typo = style_builder.build_typography(struct, theme, [])
    rounded = style_builder.build_rounded(struct, [])
    spacing = style_builder.build_spacing(struct)
    check("字体层级含 h1/h2/body", all(k in typo for k in ("h1", "h2", "body-md")), str(list(typo)[:6]))
    check("圆角档位有 sm/md", all(k in rounded for k in ("sm", "md")), str(rounded))
    solid = {k: v for k, v in rounded.items() if v != "999px"}      # pill 是胶囊，属例外
    check("圆角数值合理（≤24px，pill 除外）",
          all(int(re.sub(r"\D", "", v) or 99) <= 24 for v in solid.values()), str(rounded))
    check("字体层级随 h1→body 递减", _rem(typo.get("h1", {})) >= _rem(typo.get("body-md", {})) > 0,
          "h1=%s body=%s" % (typo.get("h1", {}).get("fontSize"), typo.get("body-md", {}).get("fontSize")))
    check("间距含 page-margin", "page-margin" in spacing, str(spacing))
    try:
        tokens = style_builder.build_tokens(palette, typo, rounded, spacing, {}, [])
        check("令牌含 colors.bg / colors.ink", "bg" in tokens.get("colors", {}) and "ink" in tokens.get("colors", {}),
              str(list(tokens.get("colors", {}).keys())[:8]))
    except Exception as exc:                            # noqa: BLE001
        tokens = {"colors": {}, "components": {}}
        check("令牌含 colors.bg / colors.ink", False, "%s: %s" % (type(exc).__name__, exc))

    sections = dict(overview="## Overview\n\n实测提取。", colors="## Colors\n\n- **accent**：主色。",
                    typography="## Typography\n\n标题与正文分层。", layout="## Layout\n\n边距与网格。",
                    elevation="## Elevation & Depth\n\n约定。", shapes="## Shapes\n\n圆角与直角。",
                    components="## Components\n\n见 front matter。",
                    dos="## Do's and Don'ts\n\n- ✅ 保持主色统一")
    style_md = style_builder.render_style_md(tokens, sections)
    check("style.md 有 front matter", style_md.startswith("---\n"), style_md[:14].replace("\n", "\\n"))
    check("style.md 含 colors/typography 段", "colors:" in style_md and "typography:" in style_md)
    check("style.md 含 Components 与 Do's and Don'ts", "## Components" in style_md and "Do's and Don'ts" in style_md)
    issues = style_builder.self_check(style_md)
    check("self_check 返回可读的问题清单", isinstance(issues, list), str(issues[:2]))

    # ── ⑥ prompt.md 1500 字符硬约束 ──
    print("\n[6] prompt.md 预算（超了会被 render 静默截掉负向词）")
    prompt_md = style_builder.build_prompt_md("测试风格", palette, {}, evidence, rounded, struct)
    check("prompt.md ≤1500 字符", len(prompt_md) <= 1500, "%d 字符" % len(prompt_md))
    check("prompt.md 含负向词段", "负向" in prompt_md)
    long_blocks = [(1, "很长的段落。" * 200), (2, "## 负向词\nwatermark, logo"), (3, "## 中英关键词\nclean")]
    fitted = style_builder._fit_prompt("# 画面语言：测试", long_blocks)
    check("超预算时被裁剪到 ≤1500", len(fitted) <= 1500, "%d 字符" % len(fitted))

    # ── ⑦ 落盘与可见性 ──
    print("\n[7] 保存成风格包（落盘 → 风格库可见）")
    bid = "sb_test_save"
    os.makedirs(style_builder._builder_dir(bid), exist_ok=True)
    pages_dir = style_builder._pages_dir(bid)
    os.makedirs(pages_dir, exist_ok=True)
    shutil.copy(png_path, os.path.join(pages_dir, "p01.png"))
    style_builder._write_state(bid, dict(
        builder_id=bid, kind="image", status="analyzed", files=[], pages=1, warnings=[], error="",
        created_at="t", result=dict(
            style_md=style_md, prompt_md=prompt_md,
            previews=["/files/style-builder/%s/source-pages/p01.png" % bid],
            palette=palette, warnings=[],
            suggested=dict(id="auto-test-style", name="测试风格", category="自动提取/待分类",
                           tags=["测试", "回归"], summary="离线回归造的包"))))

    with_kw = dict(id="auto-test-style", name="自动提取测试风格", category="自动提取/待分类",
                   tags=["测试", "自动"], summary="离线回归造的包")
    saved = style_builder.save(bid, with_kw)
    dest = os.path.join(STYLE_ROOT, "styles", "auto-test-style")
    check("save 返回 ok", saved.get("ok") is True)
    check("落了 style.md / prompt.md / meta.json / previews",
          all(os.path.exists(os.path.join(dest, name)) for name in ("style.md", "prompt.md", "meta.json", "previews")))
    meta_on_disk = json.load(io.open(os.path.join(dest, "meta.json"), encoding="utf-8"))
    check("meta.json 写入提交的展示名", meta_on_disk.get("name") == "自动提取测试风格", str(meta_on_disk)[:80])

    packs = {p["id"]: p for p in stylepacks.list_style_packs()}
    check("风格库里能看到新包", "auto-test-style" in packs)
    new_pack = packs.get("auto-test-style") or {}
    check("新包预览图计数 > 0", new_pack.get("preview_count", 0) > 0, "preview_count=%s" % new_pack.get("preview_count"))
    check("展示名来自 meta.json（非目录名）", new_pack.get("name") == "自动提取测试风格", str(new_pack.get("name")))
    check("分类/标签来自 meta.json", new_pack.get("category") == "自动提取/待分类" and new_pack.get("tags") == ["测试", "自动"],
          "%s / %s" % (new_pack.get("category"), new_pack.get("tags")))
    detail = stylepacks.get_style_pack("auto-test-style") or {}
    check("详情能取到 style_md/prompt_md/tokens",
          bool(detail.get("style_md")) and bool(detail.get("prompt_md")) and bool(detail.get("tokens")))

    # ── ⑧ meta.json 覆盖内置 META ──
    print("\n[8] meta.json 优先级（运行时新包也能有中文展示名）")
    builtin_dir = os.path.join(STYLE_ROOT, "styles", "tech-deepblue")
    os.makedirs(builtin_dir, exist_ok=True)
    with io.open(os.path.join(builtin_dir, "style.md"), "w", encoding="utf-8") as f:
        f.write("---\nname: Builtin Name\ndescription: 内置包的描述\ncolors:\n  accent: \"#3B82F6\"\n---\n\n## Overview\n")
    with io.open(os.path.join(builtin_dir, "meta.json"), "w", encoding="utf-8") as f:
        json.dump(dict(name="被 meta.json 覆盖的名字", category="覆盖分类"), f, ensure_ascii=False)
    overwritten = {p["id"]: p for p in stylepacks.list_style_packs()}.get("tech-deepblue") or {}
    check("meta.json 覆盖了内置 META 的名字", overwritten.get("name") == "被 meta.json 覆盖的名字",
          str(overwritten.get("name")))
    check("未覆盖的字段仍来自内置 META", bool(overwritten.get("summary")), str(overwritten.get("summary"))[:40])

    # ── ⑧b 并发写状态（真机踩过：固定 .tmp 名互相踩 → FileNotFoundError 整单失败）──
    print("\n[8b] 并发写状态文件（防回归）")
    import threading

    errors: list = []

    def hammer(worker: int) -> None:
        try:
            for n in range(12):
                cur = style_builder._read_state(bid)
                style_builder._write_state(bid, dict(cur, hammer="%d-%d" % (worker, n)))
        except Exception as exc:                        # noqa: BLE001
            errors.append("%s: %s" % (type(exc).__name__, exc))

    threads = [threading.Thread(target=hammer, args=(i,)) for i in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    check("6 线程 × 12 次并发写状态无异常", not errors, "；".join(errors[:2]) or "无异常")
    final = style_builder._read_state(bid)
    check("并发写后状态仍可解析且完整", bool(final) and final.get("builder_id") == bid,
          "字段数 %d" % len(final))

    # ── ⑨ 拒绝路径 ──
    print("\n[9] 拒绝与清理（不覆盖既有包、不落半成品）")
    good_state = style_builder._read_state(bid)

    def reset_analyzed():
        """save 成功会把 status 改成 saved；每个拒绝用例前要回到“可保存”状态。"""
        style_builder._write_state(bid, dict(good_state, status="analyzed", result=good_state.get("result")))

    for label, payload, key in (("重名 id 被拒", dict(id="auto-test-style"), "已存在"),
                                ("非法 id 被拒", dict(id="Bad ID!"), "不合法")):
        reset_analyzed()
        try:
            style_builder.save(bid, payload)
            check(label, False, "没有抛错")
        except ValueError as exc:
            check(label, key in str(exc), str(exc)[:60])
    reset_analyzed()
    long_result = dict(good_state.get("result") or {}, prompt_md="x" * 1600)
    style_builder._write_state(bid, dict(good_state, status="analyzed", result=long_result))
    try:
        style_builder.save(bid, dict(id="auto-too-long"))
        check("prompt.md 超 1500 被拒", False, "没有抛错")
    except ValueError as exc:
        check("prompt.md 超 1500 被拒", "1500" in str(exc), str(exc)[:60])
    check("不存在的 builder 返回 None", style_builder.response("sb_not_exist") is None)
    out = style_builder.delete(bid)
    check("delete 清掉 builder 目录", out.get("ok") is True and not os.path.isdir(style_builder._builder_dir(bid)))

    print("\n%d 项检查：%d 通过 / %d 失败" % (len(PASS) + len(FAIL), len(PASS), len(FAIL)))
    if FAIL:
        print("失败项：", "、".join(FAIL))
    shutil.rmtree(TMP, ignore_errors=True)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
