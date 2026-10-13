#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""原生渲染器回归测试：**渲染器不得伪造内容与数据**。

跑法：.venv/bin/python tools/test_native_render.py
依赖：node（server/pptx/node_modules 里的 pptxgenjs）、python-pptx

为什么需要它：render_native.mjs 曾经在版面里写死过「项目创新 / 团队管理与协同」这类栏目标题和
「以创新驱动发展…」口号，并且在要点没有数值时拿「项目 1、项目 2…」序号充当图表数据。
这两类都是**凭空造内容**，用户无法从产物里分辨真假。本测试把它们钉死为回归项。

每个用例：构造 slides → 调 node 渲染 → 用 python-pptx 打开产物 → 断言文本/图表。
"""
import json
import os
import subprocess
import sys
import tempfile

from pptx import Presentation

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RENDERER = os.path.join(ROOT, "server", "pptx", "render_native.mjs")

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print("  %s %s%s" % ("✅" if cond else "❌", name, (" — " + detail) if detail else ""))


def render(slides: list) -> str:
    """调渲染器产出一个临时 pptx，返回路径。"""
    out = os.path.join(tempfile.mkdtemp(prefix="df-native-test-"), "out.pptx")
    payload = dict(output=out, project=dict(id="t", title="渲染回归测试"),
                   tokens={"colors": {"accent": "#1E78FA"}}, slides=slides)
    proc = subprocess.run(["node", RENDERER], input=json.dumps(payload, ensure_ascii=False),
                          text=True, capture_output=True, timeout=90, check=False)
    if proc.returncode != 0 or not os.path.exists(out):
        raise RuntimeError("渲染失败：%s" % (proc.stderr or proc.stdout)[:400])
    return out


def slide_texts(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def chart_of(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_chart", False):
            return sh.chart
    return None


FORBIDDEN = ["项目创新", "团队管理与协同", "以创新驱动发展", "用创意连接更多可能", "项目 1", "项目 2"]


def main() -> int:
    print("\n[1] 对比页：没有真实栏目标签时，不得编造栏目标题与口号")
    f = render([dict(page_no=1, layout="comparison_2col", title="项目背景与痛点",
                     outline="呈现排队、缺货、库存与系统割裂四类痛点",
                     points=["校园超市高峰期平均排队时长 8.5 分钟", "学生缺货感知率达到 37%",
                             "滞销商品库存周转天数超过 45 天", "校内既有 3 个小程序互不打通"])])
    t = slide_texts(Presentation(f).slides[0])
    check("对比页无伪造栏目标题", not any(x in t for x in ["项目创新", "团队管理与协同"]))
    check("对比页无伪造口号", "以创新驱动发展" not in t)
    check("对比页保留了本页真实要点", "学生缺货感知率达到 37%" in t and "滞销商品库存周转天数超过 45 天" in t)

    print("\n[2] 对比页：要点里写了「标签：内容」时，栏目用真实标签")
    f = render([dict(page_no=1, layout="comparison_2col", title="现状与对策", outline="左右对照",
                     points=["现状：排队 8.5 分钟", "对策：预定后到店自提", "现状：缺货感知 37%", "对策：库存预警"])])
    t = slide_texts(Presentation(f).slides[0])
    check("栏目使用要点里的真实标签", ("现状" in t) and ("对策" in t))

    print("\n[3] 数据页：视觉模型识别到真实图表 → 用识别结果并标注来源（来源写实际模型名）")
    f = render([dict(page_no=1, layout="data_chart", title="运营数据", outline="试点四个月数据",
                     points=["累计注册用户 3860 人", "日均订单 412 单", "复购率 63%"],
                     visual=dict(analysis_model="qwen-vl-plus"),
                     vision=dict(chart_data=[dict(title="各校区月度订单量", labels=["1月", "2月", "3月"],
                                                  series=[dict(name="校区一", values=[800, 1200, 1600]),
                                                          dict(name="校区二", values=[600, 950, 1300])])]))])
    s = Presentation(f).slides[0]
    c = chart_of(s)
    check("识别到图表时确实画出图表", c is not None)
    if c is not None:
        cats = [str(x) for x in c.plots[0].categories]
        vals = list(c.plots[0].series[0].values)
        check("图表用的是识别出的标签与数值", cats == ["1月", "2月", "3月"] and vals == [800.0, 1200.0, 1600.0],
              "%s / %s" % (cats, vals))
        check("多系列全部画出来", len(c.plots[0].series) == 2,
              str([(sr.name, list(sr.values)) for sr in c.plots[0].series]))
    t3 = slide_texts(s)
    check("识别到的图表标题被画出", "各校区月度订单量" in t3)
    check("来源标注用实际视觉模型名", "页面图片识别（qwen-vl-plus）" in t3)
    check("来源不写死厂商", "Qwen-VL" not in t3)

    print("\n[3b] 数据页：换成自定义视觉模型时，来源标注跟着变")
    f = render([dict(page_no=1, layout="data_chart", title="运营数据", outline="试点数据",
                     points=["累计注册用户 3860 人", "日均订单 412 单"],
                     visual=dict(analysis_model="my-vision-model"),
                     vision=dict(chart_data=[dict(labels=["A", "B"], series=[dict(name="值", values=[1, 2])])]))])
    t3b = slide_texts(Presentation(f).slides[0])
    check("标注为 my-vision-model", "页面图片识别（my-vision-model）" in t3b)

    print("\n[3c] 数据页：一页识别出多组图表时，取信息量最大的一组")
    f = render([dict(page_no=1, layout="data_chart", title="多组图表", outline="识别到多组",
                     points=["要点一", "要点二"],
                     visual=dict(analysis_model="m"),
                     vision=dict(chart_data=[dict(title="小组", labels=["甲", "乙"],
                                                  series=[dict(name="s", values=[1, 2])]),
                                             dict(title="大组", labels=["甲", "乙", "丙"],
                                                  series=[dict(name="s", values=[1, 2, 3])])]))])
    t3c = slide_texts(Presentation(f).slides[0])
    check("画出的是信息量更大的那组", ("大组" in t3c) and ("小组" not in t3c), t3c.replace("\n", " / ")[:80])

    print("\n[4] 数据页：无识别结果但要点里有同单位真实数值 → 用要点数值，来源标注为要点")
    f = render([dict(page_no=1, layout="data_chart", title="运营数据", outline="试点数据",
                     points=["累计注册用户 3860 人", "活跃用户 1520 人", "校园大使 12 人"])])
    s = Presentation(f).slides[0]
    c = chart_of(s)
    check("用要点数值画出图表", c is not None)
    if c is not None:
        check("图表数值等于要点里的原文数值", list(c.plots[0].series[0].values) == [3860.0, 1520.0, 12.0],
              str(list(c.plots[0].series[0].values)))
    check("来源标注为要点原文数值", "要点原文数值" in slide_texts(s))

    print("\n[5] 数据页：要点里没有数值 → 不画图表（原来会用「项目 1/2/3」序号造假数据）")
    f = render([dict(page_no=1, layout="data_chart", title="项目概述", outline="定位与进展",
                     points=["面向高校校园零售场景构建智慧零售平台", "由学生创业团队自主设计与运营",
                             "已完成两个校区试点", "与后勤集团签订合作协议"])])
    s = Presentation(f).slides[0]
    t = slide_texts(s)
    check("没有假图表", chart_of(s) is None)
    check("没有「项目 1/2/3」占位标签", not any(x in t for x in ["项目 1", "项目 2", "项目 3"]))
    check("要点仍然完整可见（正文只出现一次，不重复）", t.count("已完成两个校区试点") == 1,
          "出现 %d 次" % t.count("已完成两个校区试点"))

    print("\n[6] 数据页：混合单位 → 只画同单位的那一组，绝不混轴")
    f = render([dict(page_no=1, layout="data_chart", title="混合指标", outline="不同量纲",
                     points=["平均排队 8.5 分钟", "缺货感知率 37%", "复购率 63%"])])
    s = Presentation(f).slides[0]
    c = chart_of(s)
    check("同单位的 2 条数值可以成图", c is not None)
    if c is not None:
        vals = list(c.plots[0].series[0].values)
        check("只取同单位那组（%），不含 8.5 分钟", vals == [37.0, 63.0], str(vals))
        check("系列名标注了单位", "%" in str(c.plots[0].series[0].name), str(c.plots[0].series[0].name))

    print("\n[7] 总结页：没有额外要点时，不写死宣传语")
    f = render([dict(page_no=1, layout="summary", title="运营数据", outline="试点四个月改善",
                     points=["注册用户 3860 人", "日均订单 412 单", "复购率 63%", "排队降至 2.1 分钟"])])
    check("总结页无伪造宣传语", "用创意连接更多可能" not in slide_texts(Presentation(f).slides[0]))

    print("\n[8] 全量兜底：任何版面都不得出现已知伪造串")
    layouts = ["cover", "section", "cards_4", "comparison_2col", "data_chart", "timeline", "summary"]
    f = render([dict(page_no=i + 1, layout=lb, title="测试页 %d" % (i + 1), outline="说明",
                     points=["要点一", "要点二", "要点三"]) for i, lb in enumerate(layouts)])
    all_text = "\n".join(slide_texts(s) for s in Presentation(f).slides)
    hit = [x for x in FORBIDDEN if x in all_text]
    check("7 个版面均无伪造串", not hit, "命中：%s" % hit if hit else "%d 页全部干净" % len(layouts))

    print("\n" + "=" * 68)
    print("通过 %d 项，失败 %d 项" % (len(PASS), len(FAIL)))
    if FAIL:
        print("失败清单：")
        for x in FAIL:
            print("  -", x)
        return 1
    print("✅ 原生渲染器「不伪造内容与数据」回归全绿")
    return 0


if __name__ == "__main__":
    sys.exit(main())
