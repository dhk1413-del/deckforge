"""Offline only: no sockets, model calls, or production database writes."""
import asyncio
import os
from pathlib import Path
import sys
import tempfile
from unittest.mock import AsyncMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from starlette.routing import Mount
from server.app import beautify, codex_engine, config, pipeline


def main():
    with tempfile.TemporaryDirectory() as tmp, patch.object(config, 'FILES_DIR', tmp), \
            patch('socket.create_connection', side_effect=AssertionError('禁止联网')), \
            patch.object(beautify.providers, '_call_edit_provider', side_effect=AssertionError('禁止调模型')), \
            patch.object(beautify.providers, '_call_image_provider', side_effect=AssertionError('禁止调模型')):
        from server.app.main import app
        root_mounts = [i for i, route in enumerate(app.routes)
                       if isinstance(route, Mount) and route.path == '/']
        for path in ('/api/codex-engine', '/api/projects/{pid}/beautify',
                     '/api/projects/{pid}/beautify/origin', '/api/projects/{pid}/beautify/preview', '/api/projects/{pid}/beautify/continue'):
            indices = [i for i, route in enumerate(app.routes) if route.path == path]
            assert indices, f'缺少 API 路由：{path}'
            assert all(i < mount for i in indices for mount in root_mounts), f'API 被静态挂载遮挡：{path}'
        print('PASS 7 路由注册顺序：Codex/美化 API 均位于根静态挂载之前')

        aggregate = [dict(provider='priced', cny=1.5, n=2, failed=0),
                     dict(provider='unpriced', cny=0, n=3, failed=2)]
        with patch.object(beautify.db, 'q', return_value=aggregate) as query:
            cost = beautify.cost_summary('cost-test')
            assert cost == dict(total_cny=1.5, image_calls=5, failed_calls=2,
                                price_configured=True, note='', by_provider=[
                                    dict(provider='priced', cny=1.5, count=2, failed=0),
                                    dict(provider='unpriced', cny=0.0, count=3, failed=2)])
            assert isinstance(cost['total_cny'], float)
            assert all(isinstance(row['cny'], float) for row in cost['by_provider'])
            assert query.call_args.args[1] == ('cost-test',)
        with patch.object(beautify.db, 'q', return_value=[]):
            cost = beautify.cost_summary('empty')
            assert cost == dict(total_cny=0.0, image_calls=0, failed_calls=0,
                                price_configured=False, by_provider=[],
                                note='未配置生图单价，金额记 0，不代表免费')
            assert isinstance(cost['total_cny'], float) and cost['note']
        print('PASS 成本汇总：混合供应商金额/调用/失败数与空记录提示（离线只读）')

        root = Path(tmp)
        pdf = root / 'source.pdf'
        with fitz.open() as doc:
            for i in range(4):
                page = doc.new_page(width=640, height=360)
                page.insert_text((40, 40), f'Original page {i + 1}')
            doc.save(pdf)
        png = root / 'source.png'
        Image.new('RGB', (640, 360), 'white').save(png)
        pptx = root / 'source.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '原稿标题\n不能改写'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text = '表格原文'
        table.cell(0, 1).text = '42%'
        prs.save(pptx)

        rendered = beautify.render_origin('pdf', str(pdf), limit=3)
        assert rendered['page_count'] == 4 and len(rendered['pages']) == 3
        first = Path(rendered['pages'][0]['path'])
        mtime = first.stat().st_mtime_ns
        assert all(p['width'] == 1600 for p in rendered['pages'])
        assert beautify.render_origin('pdf', str(pdf), limit=3) == rendered
        all_pages = beautify.render_origin('pdf', str(pdf))
        assert len(all_pages['pages']) == 4 and first.stat().st_mtime_ns == mtime
        first.unlink()
        assert Path(beautify.render_origin('pdf', str(pdf))['pages'][0]['path']).is_file()
        image = beautify.render_origin('img', str(png))
        assert image['renderer'] == 'Pillow' and image['page_count'] == 1 and image['pages'][0]['width'] == 1600
        assert beautify.render_origin('img', str(png)) == image
        from server.app import main as api
        from fastapi import HTTPException
        async def origin_endpoint():
            with patch.object(api, '_beautify_project', return_value=dict(source_path=str(pdf))):
                response = await api.beautify_origin('origin-api')
                assert response['page_count'] == 4 and len(response['pages']) == 4
                assert response['renderer'] == 'PyMuPDF'
                for page in response['pages']:
                    assert set(page) == {'page_no', 'image_url', 'width', 'height'}
                    assert page['width'] == 1600 and page['height'] == 900
                    assert page['image_url'].startswith('/files/')
                    assert (root / page['image_url'].removeprefix('/files/')).is_file()
                with patch.object(fitz.Page, 'get_pixmap', side_effect=AssertionError('缓存不应重渲')):
                    assert await api.beautify_origin('origin-api') == response
            for code in (400, 404):
                with patch.object(api, '_beautify_project', side_effect=HTTPException(code, '无原稿')):
                    try:
                        await api.beautify_origin('missing')
                    except HTTPException as exc:
                        assert exc.status_code == 400
                    else:
                        raise AssertionError('origin 缺项目或原稿应返回 400')
        async def origin_thread(fn, *args, **kwargs):
            return fn(*args, **kwargs)
        with patch('asyncio.to_thread', side_effect=origin_thread):
            asyncio.run(origin_endpoint())
        print('PASS origin 端点：全部原页、图片 URL 文件可用、重复调用不重渲、无原稿 400')
        print('PASS 1 render_origin：PDF/PNG、1600 宽、缓存复用和缺文件自愈')

        text = beautify.page_text_source(str(pptx), 1)
        assert text['source'] == 'pptx' and text['texts'] == ['原稿标题\n不能改写', '表格原文', '42%']
        assert beautify.page_text_source(str(pdf), 1)['texts'] == ['Original page 1']
        assert beautify.page_text_source(str(png), 1) == dict(title='', points=[], texts=[], source='none')
        blank = root / 'blank.pdf'
        with fitz.open() as doc:
            doc.new_page()
            doc.save(blank)
        assert beautify.page_text_source(str(blank), 1)['source'] == 'none'
        print('PASS 2 page_text_source：PPTX 表格/PDF 文字完整，纯图片不编造')

        out = root / 'editable.pptx'
        out.write_bytes(b'stale')
        items = [dict(page_no=1, image=image['pages'][0]['path'], text=text)]
        with patch.dict(os.environ, CODEX_BIN='/nonexistent'):
            result = beautify.convert_with_codex('img', items, str(out), {}, '')
        assert result['ok'] is False and result['warnings'] and not out.exists()
        assert not (beautify.directory('img') / 'work' / 'out.pptx').exists()
        print('PASS 3 Codex 不可用：明确失败，不产出或保留伪可编辑版')
        import subprocess
        with patch.object(codex_engine, 'binary', return_value='/mock/codex'), \
                patch.object(codex_engine, 'proxy_url', return_value='http://127.0.0.1:7897'), \
                patch.object(codex_engine.subprocess, 'run', return_value=subprocess.CompletedProcess([], 0, 'DONE', '')) as cli:
            result = codex_engine.run('-prompt', str(root / 'cli'), images=[str(png)], model='test')
            args, kwargs = cli.call_args
            assert result['ok'] and args[0][-2:] == ['--', '-prompt']
            assert kwargs['env']['HTTPS_PROXY'] == 'http://127.0.0.1:7897'
            cli.side_effect = subprocess.TimeoutExpired('codex', 1)
            assert codex_engine.run('test', str(root / 'cli'))['reason'] == 'timeout'
        with patch.object(codex_engine.db, 'setting_get', return_value='http://127.0.0.1:1111'), \
                patch.object(config, 'http_proxy_url', return_value='http://127.0.0.1:2222'), \
                patch('socket.create_connection') as tcp:
            tcp.side_effect = [OSError('unreachable'), tcp.return_value]
            assert codex_engine.proxy_url() == 'http://127.0.0.1:2222'
            assert tcp.call_args.kwargs['timeout'] == 1


        state = dict(phase='preview', pages=[], warnings=[], confirmed=False)
        beautify.write_state('state', state)
        assert beautify.read_state('state') == state
        beautify.write_state('state', beautify.read_state('state'))
        assert beautify.read_state('state') == state
        state.update(phase='full', confirmed=True)
        beautify.write_state('state', state)
        assert beautify.read_state('state') == state
        print('PASS 4 state.json：读写幂等、preview→full')

        with patch.object(beautify.providers, 'pick_image_provider', return_value=None):
            fallback = beautify.beautify_page('img', 1, {})
        assert fallback['ok'] is False and fallback['warning'] and fallback['cost_cny'] == 0
        assert Path(fallback['path']).read_bytes() == Path(image['pages'][0]['path']).read_bytes()
        print('PASS 5 零 key：原图直通、成本 0、ok=false + warning')

        with patch.object(beautify.providers, 'pick_image_provider', return_value=None), \
                patch.object(beautify.providers, 'generate_image', wraps=beautify.providers.generate_image) as generate, \
                patch.object(beautify.providers, 'edit_image', wraps=beautify.providers.edit_image) as fallback_edit:
            redesigned = beautify.beautify_page('img', 1, {}, 'redesign', text=text)
            assert not redesigned['ok'] and redesigned['warning'] and redesigned['cost_cny'] == 0
            assert Path(redesigned['path']).read_bytes() == Path(image['pages'][0]['path']).read_bytes()
            prompt = generate.call_args.args[0]
            import json
            for original in text['texts']:
                assert json.dumps(original, ensure_ascii=False)[1:-1] in prompt
            assert '16:9' in prompt and '真源以外' in prompt
            assert fallback_edit.call_count == 0
            empty = dict(title='', points=[], texts=[], source='none')
            missing = beautify.beautify_page('img', 1, {}, 'redesign', text=empty)
            assert not missing['ok'] and missing['cost_cny'] == 0
            assert '无文字真源' in missing['warning'] and '未编造文字' in missing['warning']
            assert generate.call_count == 1 and fallback_edit.call_count == 1
            assert empty == dict(title='', points=[], texts=[], source='none')
        print('PASS redesign 补充：逐字真源、无 key 原图降级成本 0、无文字警告且不编造')

        # Exercise real orchestration with only injected offline boundaries.
        async def offline_thread(fn, *args, **kwargs):
            return fn(*args, **kwargs)
        records = []
        project = dict(id='job', source_path=str(pdf), title='offline')
        async def phases():
            jid = await pipeline.run_beautify('job', 'style')
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] != 'running':
                    break
                await asyncio.sleep(0)
            assert pipeline.JOBS[jid]['status'] != 'running', beautify.read_state('job')
            assert pipeline.JOBS[jid]['status'] == 'done'
            assert pipeline.JOBS[jid]['done'] == pipeline.JOBS[jid]['total'] == 3
            assert convert.call_count == 0
            events = list(pipeline.job_queue(jid)._queue)
            stages = [event['stage'] for event in events if event['type'] == 'progress']
            assert stages[-1] == 'assemble' and 'codex' not in stages
            preview = beautify.read_state('job')
            assert preview['phase'] == 'preview' and len(preview['pages']) == 3
            assert 'editable' not in preview['exports'] and Path(preview['exports']['image']).exists()
            assert not (beautify.directory('job') / 'preview.pptx').exists()
            calls = edit.call_count
            jid = await pipeline.run_beautify('job', 'style', 'full')
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] != 'running':
                    break
                await asyncio.sleep(0)
            assert pipeline.JOBS[jid]['status'] != 'running', beautify.read_state('job')
            full = beautify.read_state('job')
            assert full['phase'] == 'done' and full['confirmed'] and len(full['pages']) == 4
            # 生成图片时**不再自动**转可编辑：没有 codex 阶段、没有 editable 产物、作业照样算成功
            assert convert.call_count == 0
            assert 'editable' not in full['exports']
            assert full['codex'].get('ok') is False and '尚未转换' in full['codex'].get('message', '')
            assert 'codex' not in [event.get('stage') for event in pipeline.job_queue(jid)._queue]
            assert pipeline.JOBS[jid]['status'] == 'done'
            assert pipeline.JOBS[jid]['done'] == pipeline.JOBS[jid]['total'] == 1
            assert edit.call_count == calls + 1
            assert len(Presentation(full['exports']['image']).slides) == 4
            assert all(row['kind'] != 'beautify_editable' for row in records if 'kind' in row)
            # 按需转可编辑：只有显式调用才跑 Codex；失败时如实报错且不伪造可编辑产物
            jid = await pipeline.run_beautify_editable('job', 'style')
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] != 'running':
                    break
                await asyncio.sleep(0)
            assert convert.call_count == 1 and len(convert.call_args.args[1]) == 4
            assert Path(convert.call_args.args[2]).name == 'full.pptx'
            assert 'codex' in [event.get('stage') for event in pipeline.job_queue(jid)._queue]
            assert pipeline.JOBS[jid]['status'] == 'failed'
            ondemand = beautify.read_state('job')
            assert 'editable' not in ondemand['exports']
            assert all(row['kind'] != 'beautify_editable' for row in records if 'kind' in row)
        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'q1', return_value=project), \
                patch.object(pipeline.db, 'insert', side_effect=lambda table, row: records.append(row)), \
                patch.object(pipeline.db, 'update'), patch.object(pipeline.db, 'x'), \
                patch.object(pipeline.db, 'log_usage'), \
                patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='离线', tokens={})), \
                patch.object(beautify.providers, 'pick_image_provider', return_value=None), \
                patch.object(beautify.providers, 'edit_image', wraps=beautify.providers.edit_image) as edit, \
                patch.object(beautify, 'convert_with_codex', wraps=beautify.convert_with_codex) as convert, \
                patch.dict(os.environ, CODEX_BIN='/nonexistent'):
            asyncio.run(phases())
        # 单页重生成：非破坏性（其它页文件不重写、图片版重装配、记账落页号）
        async def single_page():
            async def wait_done(job_id):
                for _ in range(20000):
                    if pipeline.JOBS[job_id]['status'] != 'running':
                        return
                    await asyncio.sleep(0)
                raise AssertionError('作业未结束')

            job_id = await pipeline.run_beautify('single', 'style', page_nos=[1, 2, 3])
            await wait_done(job_id)
            before = beautify.read_state('single')
            assert before['phase'] == 'preview' and [p['page_no'] for p in before['pages']] == [1, 2, 3]
            files = {p['page_no']: p['beauty'] for p in before['pages']}
            stamps = {n: os.stat(path).st_mtime_ns for n, path in files.items()}
            assert len(Presentation(before['exports']['image']).slides) == 3
            decks_before = sum(1 for row in records if row.get('kind') == 'beautify_image')

            job_id = await pipeline.run_beautify_page('single', 2, 'style', 'redesign')
            await wait_done(job_id)
            assert pipeline.JOBS[job_id]['status'] == 'done' and pipeline.JOBS[job_id]['done'] == 1, (pipeline.JOBS[job_id], beautify.read_state('single').get('warnings'))
            after = beautify.read_state('single')
            assert [p['page_no'] for p in after['pages']] == [1, 2, 3], '单页重生成不得改动页集合'
            assert os.stat(files[2]).st_mtime_ns >= stamps[2], '被重生成的那一页应被重写'
            assert os.stat(files[1]).st_mtime_ns == stamps[1], '其它页的文件不得被重写'
            assert os.stat(files[3]).st_mtime_ns == stamps[3], '其它页的文件不得被重写'
            assert len(Presentation(after['exports']['image']).slides) == 3, '图片版应按已有结果重新装配'
            assert sum(1 for row in records if row.get('kind') == 'beautify_image') == decks_before + 1, '导出记录应是替换而不是缺失'
            logged = [call.args for call in usage.call_args_list if call.args[1] == 2 and call.args[2] == 'beautify']
            assert logged and all(args[6] == 0 for args in logged), '账本应按页 2 记一条且零成本降级'
            assert any('单页重生成' in (args[9] or '') or '使用原图' in (args[9] or '') for args in logged), '记账明细应说明本次是单页重生成'
            for bad, why in ((4, '不在本次所选页面'),):
                try:
                    await pipeline.run_beautify_page('single', bad, 'style', 'redesign')
                except ValueError as exc:
                    assert why in str(exc)
                else:
                    raise AssertionError('未选中的页不应允许单页重生成')
            try:
                await pipeline.run_beautify_page('no-such-state', 1, 'style', 'redesign')
            except ValueError as exc:
                assert '先生成美化预览' in str(exc)
            else:
                raise AssertionError('没有美化现场时不应允许单页重生成')
        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'q1', return_value=project), \
                patch.object(pipeline.db, 'insert', side_effect=lambda table, row: records.append(row)), \
                patch.object(pipeline.db, 'update'), patch.object(pipeline.db, 'x'), \
                patch.object(pipeline.db, 'log_usage') as usage, \
                patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='离线', tokens={})), \
                patch.object(beautify.providers, 'pick_image_provider', return_value=None):
            asyncio.run(single_page())
        print('PASS 单页重生成：只重写该页、其它页与导出不动、图片版重装配、按页记账、未选页与无现场拒绝')
        # A successful full reconstruction is injected, never sent to a model.
        def successful_convert(pid, items, out_path, tokens, style_name):
            deck = Presentation()
            for item in items:
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = '\n'.join(item['text']['texts'])
            deck.save(out_path)
            return dict(ok=True, path=out_path, seconds=0, message='', warnings=[])
        async def redesign_phases():
            for phase in ('preview', 'full'):
                jid = await pipeline.run_beautify('job', 'style', phase, mode='redesign')
                for _ in range(20000):
                    if pipeline.JOBS[jid]['status'] != 'running':
                        break
                    await asyncio.sleep(0)
                assert pipeline.JOBS[jid]['status'] == 'done'
                status = beautify.read_state('job')
                assert status['mode'] == 'redesign'
                assert all(page['cost_cny'] == 0 for page in status['pages']) and status['warnings']
                assert convert.call_count == 0                    # 两个阶段都不再自动转可编辑
                assert 'editable' not in status['exports']
                if phase == 'preview':
                    assert generate.call_count == 3
                else:
                    assert generate.call_count == 4
                    assert len(Presentation(status['exports']['image']).slides) == 4
                    # 按需转可编辑：显式调用才跑（这里注入一个必定成功的转换）
                    jid2 = await pipeline.run_beautify_editable('job', 'style')
                    for _ in range(20000):
                        if pipeline.JOBS[jid2]['status'] != 'running':
                            break
                        await asyncio.sleep(0)
                    assert pipeline.JOBS[jid2]['status'] == 'done'
                    assert convert.call_count == 1
                    ondemand = beautify.read_state('job')
                    assert len(Presentation(ondemand['exports']['editable']).slides) == 4
        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'q1', return_value=project), \
                patch.object(pipeline.db, 'insert'), patch.object(pipeline.db, 'update'), \
                patch.object(pipeline.db, 'x'), patch.object(pipeline.db, 'log_usage') as usage, \
                patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='离线', tokens={})), \
                patch.object(beautify.providers, 'pick_image_provider', return_value=None), \
                patch.object(beautify.providers, 'generate_image', wraps=beautify.providers.generate_image) as generate, \
                patch.object(beautify, 'convert_with_codex', side_effect=successful_convert) as convert:
            asyncio.run(redesign_phases())
            assert usage.call_count == 4 and all(call.args[6] == 0 for call in usage.call_args_list)
        async def wait_terminal(jid):
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] in ('done', 'failed', 'cancelled'):
                    return
                await asyncio.sleep(0)
            raise AssertionError('作业没有终结')

        async def cancellation_phases():
            jid = await pipeline.run_beautify('job', 'style', 'preview')
            await wait_terminal(jid)
            before = beautify.read_state('job')['pages']
            original_render = beautify.render_origin
            def cancel_after_render(*args, **kwargs):
                result = original_render(*args, **kwargs)
                active = next(j for j in pipeline.JOBS.values()
                              if j['project_id'] == 'job' and j['status'] == 'running')
                pipeline.request_cancel(active['id'])
                return result
            with patch.object(beautify, 'render_origin', side_effect=cancel_after_render):
                jid = await pipeline.run_beautify('job', 'style', 'full')
                await wait_terminal(jid)
            cancelled = beautify.read_state('job')
            assert pipeline.JOBS[jid]['status'] == 'cancelled'
            assert cancelled['cancelled'] and cancelled['phase'] == 'preview'
            assert cancelled['pages'] == before and convert.call_count == 0
            assert 'editable' not in cancelled['exports']
            assert any(e['type'] == 'cancelled' for e in pipeline.job_queue(jid)._queue)
            calls = edit.call_count
            jid = await pipeline.run_beautify('job', 'style', 'full')
            await wait_terminal(jid)
            resumed = beautify.read_state('job')
            assert pipeline.JOBS[jid]['status'] == 'done'
            assert resumed['phase'] == 'done' and not resumed['cancelled']
            assert len(resumed['pages']) == 4 and edit.call_count == calls + 1
            assert convert.call_count == 0        # 生成图片不自动转可编辑（恢复流程同理）
            # Preview can be cancelled before any image is produced, then resumed.
            with patch.object(beautify, 'render_origin', side_effect=cancel_after_render):
                jid = await pipeline.run_beautify('job', 'style', 'preview')
                await wait_terminal(jid)
            assert beautify.read_state('job')['phase'] == 'preview'
            assert not beautify.read_state('job')['exports']
            jid = await pipeline.run_beautify('job', 'style', 'full')
            await wait_terminal(jid)
            assert pipeline.JOBS[jid]['status'] == 'done'

        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'q1', return_value=project), \
                patch.object(pipeline.db, 'insert'), patch.object(pipeline.db, 'update'), \
                patch.object(pipeline.db, 'x'), patch.object(pipeline.db, 'log_usage'), \
                patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='离线', tokens={})), \
                patch.object(beautify.providers, 'pick_image_provider', return_value=None), \
                patch.object(beautify.providers, 'edit_image', wraps=beautify.providers.edit_image) as edit, \
                patch.object(beautify, 'convert_with_codex', side_effect=successful_convert) as convert:
            asyncio.run(cancellation_phases())
        print('PASS 取消/续跑：render 后取消不调用 Codex，保留 pages/preview/图片版，零图片亦可续跑')

        async def generation_cancel():
            rows = [dict(id=f's{n}', page_no=n, title=f'原文{n}', status='pending') for n in range(1, 5)]
            job = pipeline.new_job('job', 'generate', 4)
            def generate_and_cancel(*args):
                pipeline.request_cancel(job['id'])
                return dict(ok=True, image_bytes=png.read_bytes(), provider='mock', model='offline')
            with patch.object(pipeline.providers, 'generate_image', side_effect=generate_and_cancel) as generate, \
                    patch.object(pipeline.render, 'assemble_prompt', return_value='原文'), \
                    patch.object(pipeline.db, 'update') as update:
                await pipeline._generate_worker(job['id'], project, rows, [1, 2, 3, 4], {}, 'style', '2K')
                assert generate.call_count == 1 and pipeline.JOBS[job['id']]['status'] == 'cancelled'
                assert pipeline.JOBS[job['id']]['done'] == 1
                finished = [c for c in update.call_args_list if c.args[0] == 'slides' and c.args[2].get('status') == 'done']
                assert len(finished) == 1 and finished[0].args[1] == 's1'
                assert any(e['type'] == 'cancelled' for e in pipeline.job_queue(job['id'])._queue)
        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'insert'), patch.object(pipeline.db, 'update'), \
                patch.object(pipeline.db, 'log_usage'):
            asyncio.run(generation_cancel())
        print('PASS 向导取消：线程完成页保留 done，剩余页不生图，作业 cancelled')

        from server.app import main as api
        from fastapi import HTTPException
        async def endpoints():
            with patch('asyncio.to_thread', side_effect=offline_thread), \
                    patch.object(api, '_beautify_project', return_value=project), \
                    patch.object(api.beautify, 'read_state', return_value=dict(phase='preview', mode='redesign')), \
                    patch.object(api.pipeline, 'run_beautify', new_callable=AsyncMock, return_value='mock-job') as run:
                for endpoint in (api.beautify_preview, api.beautify_continue):
                    try:
                        await endpoint('job', dict(mode='invalid'))
                    except HTTPException as exc:
                        assert exc.status_code == 400
                    else:
                        raise AssertionError('非法模式必须返回 400')
                for endpoint in (api.beautify_preview, api.beautify_continue):
                    for invalid in ([], [0], [5], '1', None, [True], [1.5]):
                        try:
                            await endpoint('job', dict(page_nos=invalid))
                        except HTTPException as exc:
                            assert exc.status_code == 400
                        else:
                            raise AssertionError(f'非法选页必须返回 400: {invalid}')
                run.assert_not_called()
                await api.beautify_preview('job', dict(style_pack_id='style', mode='redesign'))
                assert run.call_args.args[3] == 'redesign'
                await api.beautify_continue('job', {})
                assert run.call_args.args[3] == 'redesign'
            with patch.object(pipeline, 'get_job', return_value=None):
                try:
                    await api.cancel_job('missing')
                except HTTPException as exc:
                    assert exc.status_code == 404
                else:
                    raise AssertionError('不存在必须返回 404')
            for status in ('done', 'failed', 'cancelled'):
                with patch.object(pipeline, 'get_job', return_value=dict(status=status)):
                    try:
                        await api.cancel_job('terminal')
                    except HTTPException as exc:
                        assert exc.status_code == 400 and status in exc.detail
                    else:
                        raise AssertionError('终态必须返回 400')
            with patch.object(pipeline, 'get_job', return_value=dict(status='running')), \
                    patch.object(pipeline.db, 'update'):
                response = await api.cancel_job('active')
                assert response['ok'] and response['status'] == 'cancelling'
                assert pipeline.is_cancelled('active')
                pipeline.clear_cancel('active')
            assert api._job_terminal_event(dict(status='cancelled'))['type'] == 'cancelled'
        asyncio.run(endpoints())
        async def selected_pages():
            jid = await pipeline.run_beautify('selected', 'style', preview_pages=1, page_nos=[4, 2])
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] != 'running':
                    break
                await asyncio.sleep(0)
            assert pipeline.JOBS[jid]['status'] == 'done'
            assert [call.args[1] for call in edit.call_args_list] == [2]
            state = beautify.read_state('selected')
            assert state['page_nos'] == [2, 4] and [p['page_no'] for p in state['pages']] == [2]
            jid = await pipeline.run_beautify('selected', 'style', 'full', page_nos=[2, 4])
            for _ in range(20000):
                if pipeline.JOBS[jid]['status'] != 'running':
                    break
                await asyncio.sleep(0)
            assert pipeline.JOBS[jid]['status'] == 'done'
            assert [call.args[1] for call in edit.call_args_list] == [2, 4]
            assert [p['page_no'] for p in beautify.read_state('selected')['pages']] == [2, 4]
        with patch('asyncio.to_thread', side_effect=offline_thread), \
                patch.object(pipeline.db, 'q1', return_value=project), \
                patch.object(pipeline.db, 'insert'), patch.object(pipeline.db, 'update'), \
                patch.object(pipeline.db, 'x'), patch.object(pipeline.db, 'log_usage'), \
                patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='offline', tokens={})), \
                patch.object(beautify.providers, 'pick_image_provider', return_value=None), \
                patch.object(beautify, 'beautify_page', wraps=beautify.beautify_page) as edit, \
                patch.object(beautify, 'convert_with_codex', side_effect=successful_convert):
            asyncio.run(selected_pages())
        plan = beautify.prepare_content('content-check', str(pdf), [4, 2], [2])
        assert plan['page_nos'] == [2, 4]
        assert [p['title'] for p in plan['pages']] == ['Original page 2', 'Original page 4']
        assert [p['keep_material'] for p in plan['pages']] == [True, False]
        for selected, kept in (([], []), ([1], [2]), ([1], [True]), ([1], '1')):
            try:
                beautify.prepare_content('content-check', str(pdf), selected, kept)
            except ValueError:
                pass
            else:
                raise AssertionError('非法选页或保留设置应失败')
        plan['pages'][0].update(title='User confirmed title', points=['42% from user'], source='forged', texts=['forged'])
        saved = beautify.save_content('content-check', str(pdf), plan)
        assert saved['pages'][0]['source'] == 'user'
        assert saved['pages'][0]['texts'] == ['User confirmed title', '42% from user']
        assert beautify.prepare_content('content-check', str(pdf), [2], [])['pages'][0]['title'] == 'User confirmed title'
        broken = dict(saved, source_revision='outdated')
        try:
            beautify.save_content('content-check', str(pdf), broken)
        except ValueError:
            pass
        else:
            raise AssertionError('过期原稿不可覆盖确认内容')
        blank_plan = beautify.prepare_content('blank-plan', str(blank), [1], [])
        assert blank_plan['pages'][0]['texts'] == [] and blank_plan['pages'][0]['source'] == 'none'
        print('PASS 内容梳理：原文抽取/页序/保留设置校验/用户真源/保存复用/过期拒绝/空页不编造')

        async def content_endpoints():
            with patch('asyncio.to_thread', side_effect=offline_thread), patch.object(api, '_beautify_project', return_value=project), patch.object(api.db, 'q1', return_value=None):
                prepared = await api.beautify_content('content-api', dict(page_nos=[1], keep_material_page_nos=[]))
                response = api.beautify_content_save('content-api', prepared)
                assert response['pages'][0]['title'] == 'Original page 1'
            with patch.object(api, '_beautify_project', return_value=project), patch.object(api.db, 'q1', return_value={'id': 'running'}):
                for endpoint in (api.beautify_content,):
                    try:
                        await endpoint('locked', dict(page_nos=[1]))
                    except HTTPException as exc:
                        assert exc.status_code == 409
                    else:
                        raise AssertionError('运行中不可梳理内容')
                try:
                    api.beautify_content_save('locked', saved)
                except HTTPException as exc:
                    assert exc.status_code == 409
                else:
                    raise AssertionError('运行中不可保存内容')
        asyncio.run(content_endpoints())
        print('PASS 内容接口：提取/保存可用，运行中的作业返回 409')

        async def planned_pages():
            calls = []
            def record_page(pid, no, tokens, mode, size, text, **kwargs):
                # 勾了「保留原图」的页要带 keep_material=True（生图后按坐标把原图像素贴回）
                calls.append((no, mode, text['texts'], kwargs.get('keep_material')))
                return original_beautify(pid, no, tokens, mode, size, text, **kwargs)
            with patch.object(beautify, 'beautify_page', side_effect=record_page):
                jid = await pipeline.run_beautify('content-check', 'style', mode='redesign', page_nos=[2, 4])
                for _ in range(20000):
                    if pipeline.JOBS[jid]['status'] != 'running':
                        break
                    await asyncio.sleep(0)
                assert pipeline.JOBS[jid]['status'] == 'done'
            assert calls == [(2, 'img2img', ['User confirmed title', '42% from user'], True),
                             (4, 'redesign', ['Original page 4'], False)], calls
            saved['pages'][0]['points'] = ['Updated confirmed point']
            beautify.save_content('content-check', str(pdf), saved)
            try:
                await pipeline.run_beautify('content-check', 'style', 'full', mode='redesign', page_nos=[2, 4])
            except ValueError as exc:
                assert '重新生成预览' in str(exc)
            else:
                raise AssertionError('修改内容后不可复用旧预览续跑')
            try:
                await pipeline.run_beautify('content-check', 'style', mode='redesign', page_nos=[2])
            except ValueError as exc:
                assert '选页已变更' in str(exc)
            else:
                raise AssertionError('选页必须匹配已确认内容')
        original_beautify = beautify.beautify_page
        with patch('asyncio.to_thread', side_effect=offline_thread), patch.object(pipeline.db, 'q1', return_value=project), patch.object(pipeline.db, 'insert'), patch.object(pipeline.db, 'update'), patch.object(pipeline.db, 'x'), patch.object(pipeline.db, 'log_usage'), patch.object(pipeline.stylepacks, 'get_style_pack', return_value=dict(name='offline', tokens={})), patch.object(beautify.providers, 'pick_image_provider', return_value=None):
            asyncio.run(planned_pages())
        print('PASS 确认内容进入生成：逐页保留策略生效；文字修改阻断旧预览续跑；选页变更需重新确认')
        print('PASS page_nos：非法列表 400、非连续选页按原顺序预览、续跑仅生成缺页')
        print('PASS 取消端点：不存在 404、所有终态 400、运行中 cancelling、SSE 终态回放')
        print('PASS redesign 编排/端点：mode 持久化与透传、非法值 400、零成本记账、全量双导出')
        print('PASS 编排补充：预览不调用 Codex/不产可编辑版，full 重建全部页，预览页不重跑，记账页数不变')
        # —— 勾「保留原图」的页面：原 PPT 里的图片必须原样保留（坐标级，不依赖模型自觉）——
        picdir = Path(tmp) / 'pics'
        picdir.mkdir()
        block_png = picdir / 'block.png'
        Image.new('RGB', (400, 300), (200, 30, 30)).save(block_png)
        box_pptx = picdir / 'with-picture.pptx'
        prs2 = Presentation()
        prs2.slide_width, prs2.slide_height = Inches(13.333), Inches(7.5)
        slide2 = prs2.slides.add_slide(prs2.slide_layouts[6])
        slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = '标题不能被替换'
        slide2.shapes.add_picture(str(block_png), Inches(8), Inches(3.75), width=Inches(4), height=Inches(3))
        prs2.save(box_pptx)

        boxes = beautify.page_picture_boxes(str(box_pptx), 1)
        assert len(boxes) == 1, boxes
        x, y, w, h = boxes[0]
        assert abs(x - 8 / 13.333) < 0.01 and abs(y - 3.75 / 7.5) < 0.01, boxes
        assert abs(w - 4 / 13.333) < 0.01 and abs(h - 3 / 7.5) < 0.01, boxes
        assert beautify.page_picture_boxes(str(pdf), 1) == []          # PDF 源：整页位图，无页内图片元素

        origin_png = picdir / 'origin.png'
        generated_png = picdir / 'generated.png'
        Image.new('RGB', (1600, 900), (10, 20, 30)).save(origin_png)
        Image.new('RGB', (1600, 900), (250, 250, 250)).save(generated_png)
        assert beautify.preserve_pictures(origin_png, generated_png, []) == 0
        kept = beautify.preserve_pictures(origin_png, generated_png, boxes)
        assert kept == 1, kept
        with Image.open(generated_png) as out_img, Image.open(origin_png) as src_img:
            out_rgb, src_rgb = out_img.convert('RGB'), src_img.convert('RGB')
            inside = (int((x + w / 2) * 1600), int((y + h / 2) * 900))
            assert out_rgb.getpixel(inside) == src_rgb.getpixel(inside), '矩形内必须原样保留原图像素'
            assert out_rgb.getpixel((20, 20)) == (250, 250, 250), '矩形外不应被改动'

        cached = beautify.picture_boxes_cached('pic-cache', str(box_pptx), 1)
        assert cached == boxes
        assert (Path(tmp) / 'pic-cache' / 'beautify' / 'pictures.json').is_file(), '应写入按源指纹的缓存'
        print('PASS 保留原图：pptx/pdf 图片框坐标解析、矩形内逐像素贴回、矩形外不受影响、缓存落盘')

    print('通过 20 项，失败 0 项')


if __name__ == '__main__':
    main()
