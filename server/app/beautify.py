"""Preview-first beautification. Original text and honest fallback are mandatory."""
import io
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile
import threading
import time
import uuid

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from . import config, db, providers, source_pages

CONTENT_LOCK = threading.RLock()


def cost_summary(pid) -> dict:
    """Read-only aggregation of all beautify usage for this project."""
    rows = db.q(
        'SELECT provider, SUM(cny) cny, COUNT(*) n, '
        'SUM(CASE WHEN ok=0 THEN 1 ELSE 0 END) failed '
        'FROM usage_log WHERE project_id=? AND action="beautify" GROUP BY provider',
        (pid,),
    )
    by_provider = [dict(provider=row['provider'], cny=float(row['cny']),
                        count=int(row['n']), failed=int(row['failed'])) for row in rows]
    total_cny = sum((row['cny'] for row in by_provider), 0.0)
    price_configured = total_cny > 0
    return dict(total_cny=total_cny,
                image_calls=sum(row['count'] for row in by_provider),
                failed_calls=sum(row['failed'] for row in by_provider),
                price_configured=price_configured, by_provider=by_provider,
                note='' if price_configured else '未配置生图单价，金额记 0，不代表免费')


def directory(pid):
    return Path(config.FILES_DIR) / pid / 'beautify'


def read_state(pid):
    try:
        value = json.loads((directory(pid) / 'state.json').read_text(encoding='utf-8'))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


def write_state(pid, state):
    root = directory(pid)
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=root, delete=False) as f:
        json.dump(state, f, ensure_ascii=False, indent=2)
        pending = f.name
    os.replace(pending, root / 'state.json')


def fingerprint(src):
    stat = os.stat(src)
    return dict(name=os.path.basename(src), size=stat.st_size, mtime_ns=stat.st_mtime_ns)


def page_count(src):
    if str(src).lower().endswith('.pptx'):
        return len(Presentation(src).slides)
    if str(src).lower().endswith('.pdf'):
        import fitz
        with fitz.open(src) as doc:
            return doc.page_count
    return 1


def render_origin(pid, src_path, limit=None, width=1600):
    import fitz
    root = directory(pid)
    source = fingerprint(src_path)
    cache = root / 'origin' / 'cache.json'
    try:
        old = json.loads(cache.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        old = {}
    if old.get('source') != source or old.get('width') != width:
        shutil.rmtree(root / 'origin', ignore_errors=True)
        old = {}
    origin = root / 'origin'
    origin.mkdir(parents=True, exist_ok=True)
    count = page_count(src_path)
    wanted = min(count, limit) if limit is not None else count
    pages = {p['page_no']: p for p in old.get('pages', []) if Path(p['path']).is_file()}
    ext = Path(src_path).suffix.lower()
    pdf = source_pages.ensure_pdf(pid, src_path) if ext == '.pptx' else src_path
    doc = fitz.open(pdf) if ext in ('.pptx', '.pdf') else None
    try:
        for n in range(1, wanted + 1):
            if n in pages:
                continue
            path = origin / f'p{n:02d}.png'
            if doc is not None:
                page = doc[n - 1]
                pix = page.get_pixmap(matrix=fitz.Matrix(width / page.rect.width, width / page.rect.width), alpha=False)
                pix.save(str(path))
                w, h = pix.width, pix.height
            else:
                with Image.open(src_path) as image:
                    w, h = width, max(1, round(image.height * width / image.width))
                    image.convert('RGB').resize((w, h)).save(path)
            pages[n] = dict(page_no=n, path=str(path), width=w, height=h)
    finally:
        if doc is not None:
            doc.close()
    result = dict(ready=True, page_count=count, renderer='PyMuPDF' if doc is not None else 'Pillow',
                  source=source, width=width, pages=[pages[n] for n in sorted(pages)])
    cache.write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    return dict(result, pages=[pages[n] for n in range(1, wanted + 1)])


def page_text_source(src_path, page_no):
    texts = []
    source = 'none'
    if page_no < 1:
        raise ValueError('页码必须从 1 开始')
    ext = Path(src_path).suffix.lower()
    if ext == '.pptx':
        slide = Presentation(src_path).slides[page_no - 1]
        def visit(shapes):
            for shape in shapes:
                if shape.has_text_frame and shape.text.strip():
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(cell.text for row in shape.table.rows for cell in row.cells if cell.text.strip() and not cell.is_spanned)
                if shape.shape_type == 6:
                    visit(shape.shapes)
        visit(slide.shapes)
        source = 'pptx' if texts else 'none'
    elif ext == '.pdf':
        import fitz
        with fitz.open(src_path) as doc:
            texts = [line for line in doc[page_no - 1].get_text().splitlines() if line.strip()]
        source = 'pdf' if texts else 'none'
    return dict(title=texts[0] if texts else '', points=texts[1:], texts=texts, source=source)


def page_picture_boxes(src_path, page_no):
    """该页「图片元素」的归一化坐标 [x, y, w, h]（0~1，左上角为原点）。

    用途：勾了「保留原图」的页面，生图后把这些区域的**原图像素原样贴回**，
    保证「PPT 里的图片不删、不换、不被重绘」——这是坐标级保证，不依赖模型自觉。

    - pptx：python-pptx 直接给形状坐标（EMU），除以幻灯片尺寸；含组合形状内的图片
    - pdf ：PyMuPDF 的 get_image_info()，除以页面 rect
    - 其它（jpg/png 源）：整页本身就是一张位图，没有"页内图片元素"，返回空列表
    """
    if page_no < 1:
        raise ValueError('页码必须从 1 开始')
    ext = Path(src_path).suffix.lower()
    boxes = []
    if ext == '.pptx':
        deck = Presentation(src_path)
        slide_w = deck.slide_width or 1
        slide_h = deck.slide_height or 1
        slide = deck.slides[page_no - 1]

        def visit(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes)
                    continue
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                left, top = shape.left, shape.top
                width, height = shape.width, shape.height
                if None in (left, top, width, height) or width <= 0 or height <= 0:
                    continue
                x0 = max(0.0, left / slide_w)
                y0 = max(0.0, top / slide_h)
                x1 = min(1.0, (left + width) / slide_w)
                y1 = min(1.0, (top + height) / slide_h)
                if x1 - x0 <= 0.001 or y1 - y0 <= 0.001:
                    continue
                boxes.append([round(x0, 4), round(y0, 4), round(x1 - x0, 4), round(y1 - y0, 4)])

        visit(slide.shapes)
    elif ext == '.pdf':
        import fitz
        with fitz.open(src_path) as doc:
            page = doc[page_no - 1]
            rect = page.rect
            page_w = rect.width or 1
            page_h = rect.height or 1
            for info in page.get_image_info():
                rect_x0, rect_y0, rect_x1, rect_y1 = info.get('bbox') or (0, 0, 0, 0)
                x0 = max(0.0, (rect_x0 - rect.x0) / page_w)
                y0 = max(0.0, (rect_y0 - rect.y0) / page_h)
                x1 = min(1.0, (rect_x1 - rect.x0) / page_w)
                y1 = min(1.0, (rect_y1 - rect.y0) / page_h)
                if x1 - x0 <= 0.001 or y1 - y0 <= 0.001:
                    continue
                boxes.append([round(x0, 4), round(y0, 4), round(x1 - x0, 4), round(y1 - y0, 4)])
    return boxes


def picture_boxes_cached(pid, src_path, page_no):
    """按源文件指纹缓存每页的图片框，避免逐页重复解析 pptx/pdf。"""
    if not src_path or not os.path.isfile(src_path):
        return []
    root = directory(pid)
    cache_path = root / 'pictures.json'
    fingerprint = source_revision(src_path)
    data = {}
    if cache_path.is_file():
        try:
            data = json.loads(cache_path.read_text(encoding='utf-8'))
        except (OSError, ValueError):
            data = {}
    if data.get('source') != fingerprint:
        data = dict(source=fingerprint, pages={})
    pages = data.setdefault('pages', {})
    key = str(page_no)
    if key not in pages:
        try:
            pages[key] = page_picture_boxes(src_path, page_no)
        except Exception:                     # noqa: BLE001 —— 解析失败不该拖垮生图
            pages[key] = []
        root.mkdir(parents=True, exist_ok=True)
        tmp = cache_path.with_name(cache_path.name + '.pending')
        tmp.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        os.replace(tmp, cache_path)
    return pages[key]


def source_path_of(pid) -> str:
    row = db.q1('SELECT source_path FROM projects WHERE id=?', (pid,))
    return (row or {}).get('source_path') or ''


VERSION_KEEP = 5          # 每页最多保留的历史版本数（超出删最旧的，避免磁盘无限增长）


def _versions_dir(pid, page_no):
    folder = directory(pid) / 'versions' / f'p{page_no:02d}'
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def read_versions(pid, page_no) -> list:
    try:
        data = json.loads((_versions_dir(pid, page_no) / 'index.json').read_text(encoding='utf-8'))
        return [e for e in data if isinstance(e, dict) and e.get('file')]
    except (OSError, ValueError):
        return []


def write_versions(pid, page_no, entries: list) -> None:
    folder = _versions_dir(pid, page_no)
    pending = folder / 'index.json.pending'
    pending.write_text(json.dumps(entries, ensure_ascii=False), encoding='utf-8')
    os.replace(pending, folder / 'index.json')


def current_beauty(pid, page_no):
    return directory(pid) / 'beauty' / f'p{page_no:02d}.png'


def archive_version(pid, page_no, instruction=''):
    """把该页**当前**的美图归档成历史版本（重生成/恢复之前调用）。没有现存图时返回 None。"""
    target = current_beauty(pid, page_no)
    if not target.is_file():
        return None
    folder = _versions_dir(pid, page_no)
    name = time.strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:6] + '.png'
    shutil.copyfile(target, folder / name)
    entry = dict(id=name[: -len('.png')], file=name, created_at=db.now(),
                 instruction=(instruction or '').strip()[:200], size_bytes=(folder / name).stat().st_size)
    entries = [e for e in read_versions(pid, page_no) if e.get('file') != name]
    entries.insert(0, entry)
    for old in entries[VERSION_KEEP:]:
        try:
            (folder / old['file']).unlink()
        except OSError:
            pass
    write_versions(pid, page_no, entries[:VERSION_KEEP])
    return entry


def list_versions(pid, page_no) -> dict:
    """历史版本列表。当前正在使用的图**不在**列表里（它是 beauty/pNN.png），由前端单独标注。"""
    folder = _versions_dir(pid, page_no)
    entries = [dict(id=e['id'], created_at=e.get('created_at', ''), instruction=e.get('instruction', ''),
                    size_bytes=e.get('size_bytes', 0),
                    url=f'/files/{pid}/beautify/versions/p{page_no:02d}/{e["file"]}')
               for e in read_versions(pid, page_no) if (folder / e['file']).is_file()]
    live = current_beauty(pid, page_no).is_file()
    return dict(page_no=page_no, current=('/files/%s/beautify/beauty/p%02d.png' % (pid, page_no)) if live else '',
                versions=entries)


def restore_version(pid, page_no, version_id: str):
    """用某个历史版本覆盖当前图（覆盖前先把当前图也归档，保证可再切回来）。"""
    folder = _versions_dir(pid, page_no)
    entry = next((e for e in read_versions(pid, page_no) if e['id'] == version_id), None)
    if not entry:
        raise ValueError('没有这个历史版本')
    source = folder / entry['file']
    if not source.is_file():
        raise ValueError('历史版本文件已丢失')
    archive_version(pid, page_no, '恢复版本前归档')
    target = current_beauty(pid, page_no)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)
    return entry


def preserve_pictures(origin_path, target_path, boxes, pad=0.002):
    """把原图这些区域的像素**原样**贴回生成结果，返回实际处理的图片数。

    注意：只覆盖这些矩形，矩形之外的生成结果（配色/装饰/文字排版）保持模型产出。
    """
    if not boxes:
        return 0
    with Image.open(origin_path) as origin_file:
        origin = origin_file.convert('RGB')
    with Image.open(target_path) as target_file:
        generated = target_file.convert('RGB')
    if origin.size != generated.size:
        origin = origin.resize(generated.size, Image.LANCZOS)   # 同长宽比，尺寸不同时对齐
    width, height = generated.size
    applied = 0
    for box in boxes:
        x, y, w, h = box
        left = max(0, int((x - pad) * width))
        top = max(0, int((y - pad) * height))
        right = min(width, int(round((x + w + pad) * width)))
        bottom = min(height, int(round((y + h + pad) * height)))
        if right - left < 1 or bottom - top < 1:
            continue
        generated.paste(origin.crop((left, top, right, bottom)), (left, top))
        applied += 1
    if applied:
        generated.save(target_path, format='PNG')
    return applied


def source_revision(src_path):
    return hashlib.sha256(json.dumps(fingerprint(src_path), sort_keys=True).encode()).hexdigest()


def content_plan(pid, src_path):
    path = directory(pid) / 'content.json'
    try:
        plan = json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return {}
    return plan if plan.get('source_revision') == source_revision(src_path) else {}


def prepare_content(pid, src_path, page_nos, keep_material_page_nos):
    selected = validate_page_nos(page_nos, page_count(src_path), True)
    if not isinstance(keep_material_page_nos, list) or any(type(n) is not int or n not in selected for n in keep_material_page_nos):
        raise ValueError('保留原图的页码必须来自已选页面')
    revision = source_revision(src_path)
    saved = {p['page_no']: p for p in content_plan(pid, src_path).get('pages', [])}
    pages = []
    for n in selected:
        text = saved.get(n) or page_text_source(src_path, n)
        pages.append(dict(page_no=n, title=text.get('title', ''), points=text.get('points', []),
                          texts=text.get('texts', []), source=text.get('source', 'none'),
                          keep_material=n in keep_material_page_nos))
    if revision != source_revision(src_path):
        raise ValueError('原稿已变更，请重新识别页面')
    return dict(source_revision=revision, page_nos=selected, pages=pages)


def save_content(pid, src_path, payload):
    if payload.get('source_revision') != source_revision(src_path):
        raise ValueError('原稿已变更，请重新梳理内容')
    selected = validate_page_nos(payload.get('page_nos'), page_count(src_path), True)
    pages = payload.get('pages')
    if not isinstance(pages, list) or len(pages) != len(selected):
        raise ValueError('内容页必须与已选页面一致')
    normalized = []
    for n, page in zip(selected, pages):
        if not isinstance(page, dict) or type(page.get('page_no')) is not int or page['page_no'] != n:
            raise ValueError('内容页必须保持原稿页序且不得重复')
        title, points, keep = page.get('title'), page.get('points'), page.get('keep_material')
        if not isinstance(title, str) or not isinstance(points, list) or any(not isinstance(p, str) for p in points) or type(keep) is not bool:
            raise ValueError('标题、正文与保留原图设置格式不正确')
        if len(title) + sum(len(p) for p in points) > 50000:
            raise ValueError('单页文字超过 50000 字符')
        original = page_text_source(src_path, n)
        texts = ([title] if title else []) + points
        unchanged = title == original['title'] and points == original['points']
        normalized.append(dict(page_no=n, title=title, points=points, texts=texts,
                               source=original['source'] if unchanged else 'user', keep_material=keep))
    plan = dict(source_revision=source_revision(src_path), page_nos=selected, pages=normalized)
    if payload['source_revision'] != plan['source_revision']:
        raise ValueError('原稿已变更，请重新梳理内容')
    root = directory(pid)
    root.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=root, delete=False) as out:
        json.dump(plan, out, ensure_ascii=False)
        pending = out.name
    os.replace(pending, root / 'content.json')
    return plan


def content_signature(plan):
    return hashlib.sha256(json.dumps(plan, sort_keys=True, ensure_ascii=False).encode()).hexdigest() if plan else ''


def beautify_prompt(tokens):
    return ('逐字保留原图所有文字与元素位置，不得增删/改写任何文字。'
            '只换配色、质感和装饰到目标风格；不得编造数据。输出与原图同长宽比。风格令牌：'
            + json.dumps(tokens, ensure_ascii=False))


def beautify_page(pid, page_no, tokens, mode='img2img', size='2K', text=None,
                  keep_material=False, source_path='', instruction=''):
    if mode not in ('img2img', 'redesign'):
        raise ValueError('mode 只能是 img2img 或 redesign')
    root = directory(pid)
    src = root / 'origin' / f'p{page_no:02d}.png'
    target = root / 'beauty' / src.name
    target.parent.mkdir(parents=True, exist_ok=True)
    warning = ''
    if mode == 'redesign':
        if text is None:
            try:
                text = json.loads((root / 'text' / f'p{page_no:02d}.json').read_text(encoding='utf-8'))
            except (OSError, ValueError):
                text = {}
        if not (text.get('texts') or text.get('title') or text.get('points')):
            warning = f'第 {page_no} 页无文字真源，未编造文字，退回 img2img'
            mode = 'img2img'
    def decorate(prompt_text: str) -> str:
        # 附加约束必须在**调用模型之前**拼好（原先保留原图那句写在调用之后，等于没送到模型）
        extra = ''
        if keep_material:
            extra += ('；照片、插画、图表截图、图标等图片元素必须原样保留：不得重绘、替换、删除、裁剪或改变色调')
        if (instruction or '').strip():
            extra += ('；用户在上一版基础上提出的修改要求（在不违反上述逐字保留与不编造规则的前提下尽量满足）：'
                      + instruction.strip()[:400])
        return prompt_text + extra

    try:
        if mode == 'redesign':
            prompt = ('重新设计一张 16:9 演示页。以下标题与要点必须逐字保留，不得增删、改写或编造数据；'
                      '不得出现下列文字真源以外的任何文字。文字真源：'
                      + json.dumps(dict(title=text.get('title', ''), points=text.get('points', []),
                                        texts=text.get('texts', [])), ensure_ascii=False)
                      + '。风格令牌：' + json.dumps(tokens, ensure_ascii=False))
            result = providers.generate_image(decorate(prompt), size)
        else:
            prompt = beautify_prompt(tokens)
            if text and text.get('source') == 'user':
                prompt = ('保留原页图片素材与视觉结构，采用以下用户确认的全部文字替换原图文字，'
                          '不得添加其他文字或编造数据。文字真源：'
                          + json.dumps(dict(title=text.get('title', ''), points=text.get('points', []),
                                            texts=text.get('texts', [])), ensure_ascii=False)
                          + '。风格令牌：' + json.dumps(tokens, ensure_ascii=False))
            result = providers.edit_image(src.read_bytes(), decorate(prompt), size)
        if result.get('ok'):
            with Image.open(io.BytesIO(result['image_bytes'])) as image:
                image.save(target, format='PNG')
            if keep_material:
                boxes = picture_boxes_cached(pid, source_path or source_path_of(pid), page_no)
                kept = preserve_pictures(src, target, boxes)
                if kept:
                    result['preserved_pictures'] = kept
        else:
            raise ValueError(result.get('message') or '生图失败')
    except Exception as exc:
        result = locals().get('result', {})
        shutil.copyfile(src, target)
        result = dict(result, ok=False, warning=f'第 {page_no} 页使用原图：{exc}')
    if warning:
        result['warning'] = '；'.join(filter(None, (warning, result.get('warning', ''))))
    result.pop('image_bytes', None)
    return dict(result, path=str(target))


def validate_page_nos(value, count, provided=False):
    if not provided:
        return list(range(1, count + 1))
    if not isinstance(value, list) or not value or any(type(n) is not int or not 1 <= n <= count for n in value):
        raise ValueError('page_nos 必须为非空页号列表，页号须在 1..page_count 内')
    return sorted(set(value))
