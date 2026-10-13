# -*- coding: utf-8 -*-
"""PPTX 装配（python-pptx）。

图片版：每页整页图铺满 —— 秒出、同步返回、0 成本
"""
import os
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Emu

# 16:9 画布（与页面图同比例 2048×1152）
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)


def _new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])       # 6 = Blank


def _add_full_image(slide, image_path: str, w=SLIDE_W, h=SLIDE_H, left=0, top=0) -> None:
    slide.shapes.add_picture(image_path, left, top, width=w, height=h)


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def export_image_pptx(project: Dict[str, Any], rows: List[Dict[str, Any]], out_path: str) -> Dict[str, Any]:
    prs = _new_deck()
    for r in rows:
        s = _blank(prs)
        if r.get("image_path") and os.path.exists(r["image_path"]):
            _add_full_image(s, r["image_path"])
        _notes(s, "第 %s 页 · %s · 图片版（不可编辑）" % (r.get("page_no"), r.get("provider") or "-"))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return dict(pages=len(rows), editable_pages=[], passthrough_pages=[r["page_no"] for r in rows],
                size_bytes=os.path.getsize(out_path))
